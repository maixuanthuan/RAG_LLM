from pathlib import Path

from assistant.ingestion.models import Block, ParseError
from assistant.ingestion.text import normalize_text

MAX_ROWS_PER_SHEET = 50_000


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return normalize_text(str(value)).replace("\n", " ")


def parse_xlsx(path: Path) -> tuple[list[Block], list[str]]:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        raise ParseError(f"Không mở được XLSX: {e}") from e

    blocks: list[Block] = []
    warnings: list[str] = []
    for ws in wb.worksheets:
        header: list[str] | None = None
        lines: list[str] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS_PER_SHEET:
                warnings.append(f"Sheet '{ws.title}': chỉ đọc {MAX_ROWS_PER_SHEET} dòng đầu")
                break
            cells = [_cell(v) for v in row]
            if not any(cells):
                continue
            if header is None:
                header = cells
                continue
            pairs = [
                f"{h or f'Cột {j + 1}'}: {v}"
                for j, (h, v) in enumerate(zip(header, cells, strict=False))
                if v
            ]
            if pairs:
                lines.append("; ".join(pairs))
        if lines:
            blocks.append(Block(text="\n".join(lines), page=None, section=f"Sheet {ws.title}"))
    wb.close()
    return blocks, warnings
