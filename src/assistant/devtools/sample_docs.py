"""Sinh bộ tài liệu mẫu (hư cấu) để thử pipeline khi chưa có tài liệu thật."""

from pathlib import Path


def _hr_policy(path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_heading("Quy định nhân sự Công ty ABC", 0)
    d.add_heading("Chế độ nghỉ phép", 1)
    for t in [
        "Nhân viên chính thức được nghỉ phép năm 12 ngày làm việc mỗi năm, hưởng nguyên lương.",
        "Cứ mỗi 5 năm làm việc liên tục tại công ty, nhân viên được cộng thêm 1 ngày phép.",
        "Nhân viên thử việc không được hưởng phép năm; ngày nghỉ trong thời gian thử việc "
        "được tính là nghỉ không lương, trừ trường hợp có thỏa thuận khác.",
        "Đơn xin nghỉ phép phải được gửi cho quản lý trực tiếp trước ít nhất 3 ngày làm việc; "
        "nghỉ từ 5 ngày liên tục trở lên phải báo trước 10 ngày làm việc.",
    ]:
        d.add_paragraph(t)
    d.add_heading("Nghỉ việc riêng", 2)
    d.add_paragraph("Nhân viên được nghỉ 3 ngày hưởng lương khi bản thân kết hôn.")
    d.add_paragraph("Nhân viên được nghỉ 3 ngày hưởng lương khi bố mẹ, vợ hoặc chồng, con qua đời.")
    d.add_heading("Thử việc", 1)
    for t in [
        "Thời gian thử việc tối đa là 60 ngày đối với vị trí yêu cầu trình độ cao đẳng trở lên, "
        "và 30 ngày đối với vị trí yêu cầu trình độ trung cấp.",
        "Lương thử việc bằng 85% mức lương chính thức của vị trí.",
        "Khi kết thúc thử việc, quản lý trực tiếp đánh giá bằng biểu mẫu và gửi phòng Nhân sự "
        "chậm nhất 5 ngày làm việc trước ngày hết hạn thử việc.",
    ]:
        d.add_paragraph(t)
    d.add_heading("Làm thêm giờ", 1)
    for t in [
        "Làm thêm giờ phải được quản lý phê duyệt bằng văn bản trước khi thực hiện.",
        "Số giờ làm thêm tối đa là 4 giờ mỗi ngày và 40 giờ mỗi tháng.",
        "Hệ số tiền lương làm thêm giờ là 150% vào ngày thường, 200% vào ngày nghỉ hằng tuần "
        "và 300% vào ngày lễ.",
    ]:
        d.add_paragraph(t)
    d.save(str(path))


def _finance_process(path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_heading("Quy trình thanh toán và công tác phí", 0)
    d.add_heading("Tạm ứng", 1)
    d.add_paragraph(
        "Đề nghị tạm ứng phải được gửi phòng Kế toán trước ít nhất 3 ngày làm việc và có "
        "chữ ký của trưởng bộ phận."
    )
    d.add_heading("Công tác phí", 1)
    d.add_paragraph("Mức thanh toán công tác phí được quy định theo bảng dưới đây.")
    table = d.add_table(rows=4, cols=3)
    for i, row in enumerate(
        [
            ("Địa điểm công tác", "Tiền lưu trú mỗi ngày", "Phụ cấp ăn mỗi ngày"),
            ("Hà Nội và TP.HCM", "1.200.000 đồng", "150.000 đồng"),
            ("Tỉnh thành khác", "900.000 đồng", "150.000 đồng"),
            ("Nước ngoài", "Theo thỏa thuận", "Theo thỏa thuận"),
        ]
    ):
        for j, v in enumerate(row):
            table.cell(i, j).text = v
    d.add_heading("Hoàn ứng và thanh toán", 1)
    d.add_paragraph(
        "Nhân viên phải hoàn ứng trong vòng 7 ngày làm việc sau khi kết thúc công tác, "
        "kèm hóa đơn chứng từ hợp lệ."
    )
    d.add_paragraph(
        "Kế toán có quyền tạm dừng các khoản tạm ứng mới nếu nhân viên còn khoản tạm ứng quá hạn."
    )
    d.save(str(path))


def _it_policy(path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_heading("Chính sách an toàn thông tin", 0)
    d.add_heading("Mật khẩu", 1)
    d.add_paragraph(
        "Mật khẩu tài khoản công ty phải dài tối thiểu 12 ký tự, gồm chữ hoa, chữ thường và số."
    )
    d.add_paragraph(
        "Mật khẩu phải được đổi mỗi 90 ngày và không được trùng với 5 mật khẩu gần nhất."
    )
    d.add_heading("Làm việc từ xa", 1)
    d.add_paragraph(
        "Khi làm việc từ xa, nhân viên bắt buộc kết nối VPN của công ty và bật xác thực hai yếu tố."
    )
    d.add_heading("Thiết bị", 1)
    d.add_paragraph("Thiết bị cá nhân dùng cho công việc phải được đăng ký với phòng IT.")
    d.add_paragraph(
        "Khi mất laptop hoặc điện thoại công ty, nhân viên phải báo cho IT helpdesk "
        "(it-helpdesk@abc.example, số nội bộ 1234) trong vòng 24 giờ."
    )
    d.save(str(path))


def _allowances(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Phụ cấp"
    ws.append(
        [
            "Chức danh",
            "Phụ cấp xăng xe (đồng/tháng)",
            "Phụ cấp điện thoại (đồng/tháng)",
            "Phụ cấp ăn trưa (đồng/tháng)",
        ]
    )
    ws.append(["Nhân viên", "500.000", "200.000", "730.000"])
    ws.append(["Trưởng nhóm", "800.000", "300.000", "730.000"])
    ws.append(["Trưởng phòng", "1.200.000", "500.000", "730.000"])
    ws.append(["Giám đốc", "2.000.000", "1.000.000", "730.000"])
    wb.save(str(path))


_ONBOARDING = """HƯỚNG DẪN NHÂN VIÊN MỚI

1. Ngày đầu tiên
Nhân viên mới đến công ty lúc 8:30. Laptop được nhận tại phòng IT, tầng 3.
Quản lý trực tiếp sẽ giới thiệu nhóm và hướng dẫn chỗ ngồi.

2. Tài khoản
Email công ty được tạo trước ngày đi làm. Mật khẩu tạm thời được gửi qua số điện thoại
đã đăng ký và phải đổi ngay trong lần đăng nhập đầu tiên.

3. Giờ làm việc
Giờ làm việc từ 8:30 đến 17:30, nghỉ trưa từ 12:00 đến 13:00, từ thứ Hai đến thứ Sáu.

4. Liên hệ
Mọi thắc mắc về nhân sự gửi hr@abc.example.
"""


def make_sample_docs(out_dir: Path) -> list[Path]:
    """Tạo cấu trúc: HR/, Finance/, IT/ (tên thư mục = phòng ban). Trả về các file đã tạo."""
    out_dir = Path(out_dir)
    for sub in ("HR", "Finance", "IT"):
        (out_dir / sub).mkdir(parents=True, exist_ok=True)
    files = [
        (out_dir / "HR" / "Quy_dinh_nhan_su.docx", _hr_policy),
        (out_dir / "HR" / "Bang_phu_cap.xlsx", _allowances),
        (out_dir / "Finance" / "Quy_trinh_thanh_toan.docx", _finance_process),
        (out_dir / "IT" / "Chinh_sach_CNTT.docx", _it_policy),
    ]
    for path, fn in files:
        fn(path)
    onboarding = out_dir / "HR" / "Huong_dan_onboarding.txt"
    onboarding.write_text(_ONBOARDING, encoding="utf-8")
    return [p for p, _ in files] + [onboarding]
