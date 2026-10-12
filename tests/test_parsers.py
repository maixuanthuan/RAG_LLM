import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from assistant.ingestion.models import ParseError
from assistant.ingestion.parsers import parse_file


def test_txt_sections(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text(
        "QUY ĐỊNH CHUNG\n\n1. Phạm vi\nÁp dụng cho mọi nhân viên.\n\n2. Giờ làm\n8:30 đến 17:30\n",
        "utf-8",
    )
    doc = parse_file(f)
    assert [(b.section, b.text) for b in doc.blocks] == [
        ("1. Phạm vi", "Áp dụng cho mọi nhân viên."),
        ("2. Giờ làm", "8:30 đến 17:30"),
    ]


def test_markdown_headings(tmp_path):
    f = tmp_path / "a.md"
    f.write_text("# Tiêu đề\nnội dung\n## Mục con\nnội dung con\n", "utf-8")
    doc = parse_file(f)
    assert [b.section for b in doc.blocks] == ["Tiêu đề", "Mục con"]


def test_txt_legacy_encoding(tmp_path):
    f = tmp_path / "a.txt"
    f.write_bytes("Nghỉ phép năm".encode("utf-16"))
    assert parse_file(f).blocks[0].text == "Nghỉ phép năm"


def test_docx_headings_breadcrumb_and_table(tmp_path):
    f = tmp_path / "a.docx"
    d = Document()
    d.add_heading("Quy định", 0)
    d.add_heading("Nghỉ phép", 1)
    d.add_paragraph("Được nghỉ 12 ngày.")
    d.add_heading("Nghỉ cưới", 2)
    d.add_paragraph("Được nghỉ 3 ngày.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Chức danh", "Phụ cấp"
    t.cell(1, 0).text, t.cell(1, 1).text = "Trưởng phòng", "1.000.000"
    d.save(str(f))
    blocks = parse_file(f).blocks
    assert blocks[0].section == "Quy định > Nghỉ phép" and blocks[0].text == "Được nghỉ 12 ngày."
    assert blocks[1].section == "Quy định > Nghỉ phép > Nghỉ cưới"
    assert "Chức danh: Trưởng phòng; Phụ cấp: 1.000.000" in blocks[1].text


def test_docx_without_styles_uses_heuristic(tmp_path):
    f = tmp_path / "a.docx"
    d = Document()
    d.add_paragraph("Điều 1. Phạm vi")
    d.add_paragraph("Áp dụng cho tất cả.")
    d.save(str(f))
    blocks = parse_file(f).blocks
    assert blocks[0].section == "Điều 1. Phạm vi"


def test_xlsx_rows_become_labelled_lines(tmp_path):
    f = tmp_path / "a.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Phụ cấp"
    ws.append(["Chức danh", "Xăng xe"])
    ws.append(["Nhân viên", 500000])
    ws.append([None, None])
    ws.append(["Trưởng phòng", 1200000.0])
    wb.save(str(f))
    blocks = parse_file(f).blocks
    assert blocks[0].section == "Sheet Phụ cấp"
    assert blocks[0].text.split("\n") == [
        "Chức danh: Nhân viên; Xăng xe: 500000",
        "Chức danh: Trưởng phòng; Xăng xe: 1200000",
    ]


def make_pdf(path, pages):
    c = canvas.Canvas(str(path))
    for i, lines in enumerate(pages, start=1):
        y = 800
        for line in ["CONG TY ABC - NOI BO", *lines, f"Trang {i}/{len(pages)}"]:
            c.drawString(50, y, line)
            y -= 20
        c.showPage()
    c.save()


def test_pdf_pages_headings_and_header_footer_removed(tmp_path):
    f = tmp_path / "a.pdf"
    make_pdf(
        f,
        [
            ["Dieu 1. Pham vi", "Ap dung cho moi nhan vien", "trong cong ty."],
            ["Dieu 2. Nghi phep", "Nhan vien duoc nghi 12 ngay."],
            ["Tiep tuc noi dung dieu 2."],
        ],
    )
    doc = parse_file(f)
    assert doc.pages == 3
    sections = [(b.page, b.section) for b in doc.blocks]
    assert sections == [(1, "Dieu 1. Pham vi"), (2, "Dieu 2. Nghi phep"), (3, "Dieu 2. Nghi phep")]
    all_text = " ".join(b.text for b in doc.blocks)
    assert "CONG TY ABC" not in all_text and "Trang 2" not in all_text
    assert "Ap dung cho moi nhan vien trong cong ty." in all_text  # nối dòng bị ngắt


def test_pdf_scanned_page_warns(tmp_path):
    f = tmp_path / "a.pdf"
    c = canvas.Canvas(str(f))
    c.showPage()
    c.save()
    doc = parse_file(f)
    assert doc.blocks == []
    assert any("OCR" in w for w in doc.warnings)


def test_corrupt_and_unsupported(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a docx")
    with pytest.raises(ParseError):
        parse_file(bad)
    other = tmp_path / "a.pptx"
    other.write_bytes(b"x")
    with pytest.raises(ParseError):
        parse_file(other)
